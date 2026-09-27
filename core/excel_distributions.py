"""
Hojas de Excel para la Validacion de las Conversiones Estadisticas
===================================================================
Agrega al libro exportado una hoja por distribucion (V1 Uniforme, V2 Normal,
V3 Erlang, V4 Poisson, V5 Binomial) y una hoja resumen (V0). Todo se calcula
con FORMULAS NATIVAS de Excel a partir de la columna R_i de la hoja
"Secuencia y Calculos":

  Paso 2  Conversion:   =A+(B-A)*r, INV.NORM, INV.GAMMA, COINCIDIR sobre la
                        tabla acumulada, SUMA / PRODUCTO / CONTAR.SI de grupos...
  Paso 3  Conteo:       CONTAR.SI.CONJUNTO(x; "<"&valor) ...
  Paso 4  Validacion:   DISTR.NORM.N, POISSON.DIST, DISTR.BINOM.N ... y la regla
                        |P_sim - P_teo| <= Z_(a/2)*RAIZ(P_teo(1-P_teo)/n)

El nivel de significancia, los operadores (< / <=, > / >=) y los valores de
cada caso son editables (celdas amarillas) y la hoja se recalcula sola.
Los parametros que definen la ESTRUCTURA de la hoja (k en la convolucion de
Erlang, n de la Binomial, lambda de la tabla inversa de Poisson) quedan fijos:
para cambiarlos se ajustan en la app y se vuelve a exportar.
"""

from typing import Any, Dict, List, Optional

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .distributions import (
    CASE_ORDER,
    DIST_META,
    DIST_ORDER,
    METHOD_LABELS,
    _param_errors,
    conversion_formula,
    fmt,
    normalize_config,
    poisson_table_size,
)
from .excel_exporter import (
    COLOR_AMBER_LIGHT,
    COLOR_GREEN_DARK,
    COLOR_GREEN_LIGHT,
    COLOR_PURPLE,
    COLOR_PURPLE_DARK,
    COLOR_TEXT,
    COLOR_WHITE,
    COLOR_ZEBRA,
    _alpha_input_row,
    _kv_row,
    _objective_row,
    _section_header,
    _table_header,
    _title_banner,
    create_thin_border,
)

SHEET_NAMES = {
    "uniforme": "V1 Uniforme",
    "normal": "V2 Normal",
    "erlang": "V3 Erlang",
    "poisson": "V4 Poisson",
    "binomial": "V5 Binomial",
}
SUMMARY_SHEET = "V0 Resumen Distribuciones"

_SUBSCRIPTS = str.maketrans("0123456789", "₀₁₂₃₄₅₆₇₈₉")
COLOR_GREY_LIGHT = "F3F4F6"


def _r(i: int) -> str:
    return "r" + str(i).translate(_SUBSCRIPTS)


# -------------------------------------------------------------------------
# Parametros: (clave, etiqueta, editable)
# -------------------------------------------------------------------------
def _param_rows(dist: str, method: str):
    if dist == "uniforme":
        return [("a", "A  (límite inferior)", True), ("b", "B  (límite superior)", True)]
    if dist == "normal":
        return [("mu", "μ  (media)", True), ("sigma", "σ  (desviación estándar)", True)]
    if dist == "erlang":
        return [("k", "k  (número de fases)", method == "inversa"), ("lam", "λ  (tasa de cada fase; media = k/λ)", True)]
    if dist == "poisson":
        return [("lam", "λ  (tasa media)", method == "multiplicativo")]
    return [("n", "n  (número de ensayos)", False), ("p", "p  (probabilidad de éxito)", True)]


def _param_row(ws, row, label, value, tb, editable):
    lbl = ws.cell(row=row, column=1, value=label)
    lbl.font = Font(name="Arial", size=10, bold=True)
    lbl.alignment = Alignment(horizontal="left", vertical="center")
    val = ws.cell(row=row, column=2, value=value)
    val.font = Font(name="Arial", size=11, bold=True, color=COLOR_GREEN_DARK if editable else COLOR_TEXT)
    val.fill = PatternFill(start_color=COLOR_AMBER_LIGHT if editable else COLOR_GREY_LIGHT, fill_type="solid")
    val.alignment = Alignment(horizontal="center", vertical="center")
    val.border = tb
    note = ws.cell(row=row, column=3, value=(
        "← Editable" if editable else "Fijo: define la estructura de la hoja (cámbielo en la app y vuelva a exportar)"))
    note.font = Font(name="Arial", size=9, italic=True, color=COLOR_TEXT)
    return f"$B${row}"


def _style_cell(c, tb, bold=False, fmt_=None, fill=None, align="center", color=COLOR_TEXT):
    c.font = Font(name="Arial", size=10, bold=bold, color=color)
    c.alignment = Alignment(horizontal=align, vertical="center")
    c.border = tb
    if fill:
        c.fill = PatternFill(start_color=fill, fill_type="solid")
    if fmt_:
        c.number_format = fmt_


# -------------------------------------------------------------------------
# Probabilidad teorica como formula de Excel
# -------------------------------------------------------------------------
def _cdf_cont(dist: str, P: Dict[str, str], t: str) -> str:
    if dist == "uniforme":
        return f"MAX(0,MIN(1,({t}-{P['a']})/({P['b']}-{P['a']})))"
    if dist == "normal":
        return f"_xlfn.NORM.DIST({t},{P['mu']},{P['sigma']},TRUE)"
    # Erlang: F(x) = 1 - sum_{n=0}^{k-1} e^(-lam x)(lam x)^n/n! = 1 - POISSON.DIST(k-1; lam*x; VERDADERO)
    return f"IF({t}<=0,0,1-_xlfn.POISSON.DIST({P['k']}-1,{P['lam']}*{t},TRUE))"


def _cdf_disc(dist: str, P: Dict[str, str], t: str) -> str:
    if dist == "poisson":
        return f"IF({t}<0,0,_xlfn.POISSON.DIST({t},{P['lam']},TRUE))"
    return f"IF({t}<0,0,IF({t}>={P['n']},1,_xlfn.BINOM.DIST({t},{P['n']},{P['p']},TRUE)))"


def _theo_formula(dist: str, P: Dict[str, str], case_id: str, op_cell: str, x_cell: str, b_cell: str) -> str:
    discrete = DIST_META[dist]["discrete"]
    if not discrete:
        F = lambda t: _cdf_cont(dist, P, t)
        if case_id == "menor":
            return f"={F(x_cell)}"
        if case_id == "mayor":
            return f"=1-{F(x_cell)}"
        return f"=MAX(0,{F(b_cell)}-{F(x_cell)})"
    Fd = lambda t: _cdf_disc(dist, P, t)
    below = lambda c: f"(-INT(-{c})-1)"   # P(X < c) = F(techo(c) - 1)
    upto = lambda c: f"INT({c})"          # P(X <= c) = F(piso(c))
    if case_id == "menor":
        return f'=IF({op_cell}="<",{Fd(below(x_cell))},{Fd(upto(x_cell))})'
    if case_id == "mayor":
        return f'=IF({op_cell}=">",1-{Fd(upto(x_cell))},1-{Fd(below(x_cell))})'
    return f"=MAX(0,{Fd(upto(b_cell))}-{Fd(below(x_cell))})"


def _theo_moments(dist: str, P: Dict[str, str]):
    if dist == "uniforme":
        return f"=({P['a']}+{P['b']})/2", f"=({P['b']}-{P['a']})^2/12", "(A + B)/2", "(B − A)²/12"
    if dist == "normal":
        return f"={P['mu']}", f"={P['sigma']}^2", "μ", "σ²"
    if dist == "erlang":
        return f"={P['k']}/{P['lam']}", f"={P['k']}/{P['lam']}^2", "k/λ", "k/λ²"
    if dist == "poisson":
        return f"={P['lam']}", f"={P['lam']}", "λ", "λ"
    return f"={P['n']}*{P['p']}", f"={P['n']}*{P['p']}*(1-{P['p']})", "n·p", "n·p·(1 − p)"


# -------------------------------------------------------------------------
# Hoja de una distribucion
# -------------------------------------------------------------------------
def _build_dist_sheet(wb, order: int, dist: str, c: Dict[str, Any], ws1_title: str,
                      ri_col: str, start: int, n: int, alpha: float) -> Dict[str, Any]:
    meta = DIST_META[dist]
    method = c["method"]
    p = c["params"]
    ws = wb.create_sheet(title=SHEET_NAMES[dist])
    ws.views.sheetView[0].showGridLines = True
    tb = create_thin_border()

    _title_banner(ws, f"VALIDACIÓN DE LA CONVERSIÓN ESTADÍSTICA — {meta['name'].upper()}", span_cols=10)
    _objective_row(ws, 3, (
        f"Paso 1: se toman los rᵢ de la hoja '{ws1_title}'.  Paso 2: se convierten a la distribución con la fórmula del "
        "método.  Paso 3: se CUENTA cuántos xᵢ son menores que, mayores que o están entre dos valores "
        "(CONTAR.SI.CONJUNTO) y se divide entre n.  Paso 4: se compara con la probabilidad teórica de la distribución; la "
        "conversión queda VALIDADA si |P_sim − P_teo| ≤ Z_(α/2)·√(P_teo·(1 − P_teo)/n)."), span_cols=10)
    ws.row_dimensions[3].height = 60

    info = {"sheet": SHEET_NAMES[dist], "name": meta["name"], "method": METHOD_LABELS[(dist, method)],
            "params": None, "n_cell": None, "ok_cell": None, "error": None}

    row = 5
    errs = _param_errors(dist, p)
    if errs:
        cell = ws.cell(row=row, column=1, value="No se pudo construir la validación: " + " ".join(errs))
        cell.font = Font(name="Arial", size=11, bold=True, color="B91C1C")
        info["error"] = " ".join(errs)
        ws.column_dimensions["A"].width = 60
        return info

    rows_def = _param_rows(dist, method)
    _section_header(ws, row, "PARÁMETROS  (celdas amarillas = editables; la hoja se recalcula sola)")
    row += 1
    alpha_cell = _alpha_input_row(ws, row, alpha, tb)
    row += 1
    P: Dict[str, str] = {}
    for key, label, editable in rows_def:
        val = int(p[key]) if key in ("k", "n") else p[key]
        P[key] = _param_row(ws, row, label, val, tb, editable)
        row += 1
    formula, _formula_p, r_per_value = conversion_formula(dist, method, p)
    lbl = ws.cell(row=row, column=1, value="Método de conversión (Paso 2)")
    lbl.font = Font(name="Arial", size=10, bold=True)
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=10)
    mcell = ws.cell(row=row, column=2, value=f"{METHOD_LABELS[(dist, method)]}:   {formula}   ·   {r_per_value} rᵢ por valor")
    _style_cell(mcell, tb, bold=True, align="left", fill=COLOR_GREEN_LIGHT, color=COLOR_GREEN_DARK)
    row += 1
    z_cell = _kv_row(ws, row, "Z_(α/2)  =  NORM.S.INV(1 − α/2)", f"=_xlfn.NORM.S.INV(1-{alpha_cell}/2)", tb, num_fmt="0.0000")
    row += 1
    lim_cell = None
    if dist == "poisson" and method == "multiplicativo":
        lim_cell = _kv_row(ws, row, "e^(−λ)  (límite del producto)", f"=EXP(-{P['lam']})", tb, num_fmt="0.000000")
        row += 1
    n_row = row
    row += 2

    # ---- Filas fijas de las secciones siguientes ----
    res_header = row                  # titulo de seccion
    res_th = row + 1                  # encabezado de tabla
    case_rows = {cid: row + 2 + i for i, cid in enumerate(CASE_ORDER)}
    decision_row = row + 6
    ok_row = row + 7
    mom_header = row + 9
    mom_rows = [row + 10 + i for i in range(4)]
    conv_header = row + 15
    conv_th = row + 16
    data0 = row + 17

    # ---- Estructura de la tabla de conversion ----
    R_abs = f"'{ws1_title}'!${ri_col}${start}:${ri_col}${start + n - 1}"
    rref = lambda i: f"'{ws1_title}'!{ri_col}{start + i - 1}"
    group = None
    if (dist, method) == ("normal", "tlc"):
        group = 12
    elif (dist, method) == ("erlang", "convolucion"):
        group = int(p["k"])
    elif (dist, method) == ("binomial", "bernoulli"):
        group = int(p["n"])
    lookup = (dist, method) in (("poisson", "inversa"), ("binomial", "inversa"))
    mult = (dist, method) == ("poisson", "multiplicativo")

    if group:
        n_rows = n // group
        aux_header = {"normal": "Σ rⱼ  (12 rᵢ)", "erlang": f"Π rⱼ  ({group} rᵢ)", "binomial": "Éxitos = #{rⱼ < p}"}[dist]
        headers = ["i", "rᵢ usados", aux_header, "xᵢ convertido"]
        x_col = 4
    elif mult:
        n_rows = n
        headers = ["i", "rᵢ", "Nº de factores", "Π rⱼ acumulado", "xᵢ (cuando Π < e^(−λ))"]
        x_col = 5
    else:
        n_rows = n
        headers = ["i", "rᵢ", "xᵢ convertido"]
        x_col = 3
    first_ind_col = x_col + 1
    last_data = data0 + max(n_rows, 1) - 1
    xl = chr(ord("A") + x_col - 1)
    xR = f"${xl}${data0}:${xl}${last_data}"

    # ---- n (valores convertidos) ----
    n_cell = _kv_row(ws, n_row, "n  (valores convertidos)  =  CONTAR(xᵢ)", f"=COUNT({xR})", tb, num_fmt="0", highlight=True)
    info["n_cell"] = n_cell

    # ---- Pasos 3 y 4: tabla de resultados ----
    _section_header(ws, res_header, "RESULTADOS DE LA VALIDACIÓN  (Paso 3: conteo  ·  Paso 4: comparación con la probabilidad teórica)")
    _table_header(ws, res_th, ["Caso", "Operador", "Valor x  (o a)", "Valor b", "Conteo  CONTAR.SI.CONJUNTO",
                               "P simulada = conteo / n", "P teórica (fórmula)", "|Diferencia|",
                               "Margen = Z·√(P(1−P)/n)", "Decisión"], tb)
    ws.row_dimensions[res_th].height = 32
    labels = {"menor": "Menor que", "mayor": "Mayor que", "rango": "Entre a y b (a ≤ X ≤ b)"}
    dv_menor = DataValidation(type="list", formula1='"<,<="', allow_blank=False)
    dv_mayor = DataValidation(type="list", formula1='">,>="', allow_blank=False)
    ws.add_data_validation(dv_menor)
    ws.add_data_validation(dv_mayor)
    case_cells = {}
    for cid in CASE_ORDER:
        r = case_rows[cid]
        cs = c["cases"][cid]
        _style_cell(ws.cell(row=r, column=1, value=labels[cid]), tb, bold=True, align="left")
        if cid == "rango":
            op_c = ws.cell(row=r, column=2, value="entre")
            _style_cell(op_c, tb, fill=COLOR_GREY_LIGHT)
            a_c = ws.cell(row=r, column=3, value=cs["a"])
            b_c = ws.cell(row=r, column=4, value=cs["b"])
            _style_cell(a_c, tb, bold=True, fill=COLOR_AMBER_LIGHT, color=COLOR_GREEN_DARK)
            _style_cell(b_c, tb, bold=True, fill=COLOR_AMBER_LIGHT, color=COLOR_GREEN_DARK)
            count_f = f'=COUNTIFS({xR},">="&$C${r},{xR},"<="&$D${r})'
        else:
            op_c = ws.cell(row=r, column=2, value=cs["op"])
            _style_cell(op_c, tb, bold=True, fill=COLOR_AMBER_LIGHT, color=COLOR_GREEN_DARK)
            (dv_menor if cid == "menor" else dv_mayor).add(op_c)
            x_c = ws.cell(row=r, column=3, value=cs["x"])
            _style_cell(x_c, tb, bold=True, fill=COLOR_AMBER_LIGHT, color=COLOR_GREEN_DARK)
            _style_cell(ws.cell(row=r, column=4, value="—"), tb, fill=COLOR_GREY_LIGHT)
            count_f = f"=COUNTIFS({xR},$B${r}&$C${r})"
        case_cells[cid] = (f"$B${r}", f"$C${r}", f"$D${r}")
        _style_cell(ws.cell(row=r, column=5, value=count_f), tb, bold=True, fmt_="0")
        _style_cell(ws.cell(row=r, column=6, value=f'=IF({n_cell}=0,0,E{r}/{n_cell})'), tb, bold=True, fmt_="0.000000",
                    fill=COLOR_GREEN_LIGHT, color=COLOR_GREEN_DARK)
        _style_cell(ws.cell(row=r, column=7, value=_theo_formula(dist, P, cid, f"$B${r}", f"$C${r}", f"$D${r}")),
                    tb, bold=True, fmt_="0.000000", fill=COLOR_GREEN_LIGHT, color=COLOR_GREEN_DARK)
        _style_cell(ws.cell(row=r, column=8, value=f"=ABS(F{r}-G{r})"), tb, fmt_="0.000000")
        _style_cell(ws.cell(row=r, column=9, value=f"=IF({n_cell}=0,0,{z_cell}*SQRT(G{r}*(1-G{r})/{n_cell}))"), tb, fmt_="0.000000")
        dec = ws.cell(row=r, column=10, value=f'=IF(AND({n_cell}>=10,H{r}<=I{r}+1E-12),"✓ VALIDADA","✗ NO VALIDADA")')
        _style_cell(dec, tb, bold=True, fill=COLOR_PURPLE, color="FFFFFF")

    r1, r3 = case_rows["menor"], case_rows["rango"]
    lbl = ws.cell(row=decision_row, column=1, value="DECISIÓN FINAL:")
    lbl.font = Font(name="Arial", size=11, bold=True)
    ws.merge_cells(start_row=decision_row, start_column=2, end_row=decision_row, end_column=7)
    dec_all = ws.cell(row=decision_row, column=2, value=(
        f'=IF(B{ok_row}=3,"✓ CONVERSIÓN VALIDADA: las probabilidades simuladas coinciden con las teóricas en los 3 casos",'
        f'"✗ CONVERSIÓN NO VALIDADA: "&(3-B{ok_row})&" caso(s) fuera del margen de error")'))
    dec_all.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    dec_all.fill = PatternFill(start_color=COLOR_PURPLE, fill_type="solid")
    dec_all.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[decision_row].height = 20
    ok_cell = _kv_row(ws, ok_row, "Casos validados (de 3)",
                      f'=SUMPRODUCT(--(J{r1}:J{r3}="✓ VALIDADA"))', tb, num_fmt="0")
    info["ok_cell"] = ok_cell

    # ---- Momentos ----
    _section_header(ws, mom_header, "COMPARACIÓN DE MOMENTOS (simulado vs. teórico)")
    mt, vt, mf, vf = _theo_moments(dist, P)
    _kv_row(ws, mom_rows[0], "Media simulada  x̄ = PROMEDIO(xᵢ)", f"=IF({n_cell}=0,\"\",AVERAGE({xR}))", tb, num_fmt="0.000000")
    _kv_row(ws, mom_rows[1], f"Media teórica  E[X] = {mf}", mt, tb, num_fmt="0.000000")
    _kv_row(ws, mom_rows[2], "Varianza simulada  s² = VAR(xᵢ)", f"=IF({n_cell}<2,\"\",VAR({xR}))", tb, num_fmt="0.000000")
    _kv_row(ws, mom_rows[3], f"Varianza teórica  V[X] = {vf}", vt, tb, num_fmt="0.000000")

    # ---- Paso 2: tabla de conversion ----
    _section_header(ws, conv_header, f"TABLA DE CONVERSIÓN (Paso 2)  —  {METHOD_LABELS[(dist, method)]}:  {formula}")
    # Encabezados dinamicos: muestran la condicion vigente de cada caso (1 = cumple)
    rm, rM, rr_ = case_rows["menor"], case_rows["mayor"], case_rows["rango"]
    ind_headers = [
        f'="¿x "&$B${rm}&" "&$C${rm}&"?  (1 = sí)"',
        f'="¿x "&$B${rM}&" "&$C${rM}&"?  (1 = sí)"',
        f'="¿"&$C${rr_}&" ≤ x ≤ "&$D${rr_}&"?  (1 = sí)"',
    ]
    _table_header(ws, conv_th, headers + ind_headers, tb)
    ws.row_dimensions[conv_th].height = 30

    # Tabla acumulada para la transformada inversa discreta
    if lookup:
        T = poisson_table_size(p["lam"]) if dist == "poisson" else int(p["n"]) + 1
        tcol = first_ind_col + len(CASE_ORDER) + 1   # una columna en blanco de separacion
        tL = chr(ord("A") + tcol - 1)
        tLB = chr(ord("A") + tcol)       # limite inferior F(t-1)
        tF = chr(ord("A") + tcol + 1)    # F(t)
        _table_header(ws, conv_th, ["t", "F(t − 1)  (límite inf.)", "F(t)"], tb, start_col=tcol)
        for j in range(T):
            rr = data0 + j
            _style_cell(ws.cell(row=rr, column=tcol, value=j), tb)
            if dist == "poisson":
                F = f"=_xlfn.POISSON.DIST({tL}{rr},{P['lam']},TRUE)"
            else:
                F = f"=_xlfn.BINOM.DIST({tL}{rr},{P['n']},{P['p']},TRUE)"
            _style_cell(ws.cell(row=rr, column=tcol + 2, value=F), tb, fmt_="0.000000")
            lb = 0 if j == 0 else f"={tF}{rr - 1}"
            _style_cell(ws.cell(row=rr, column=tcol + 1, value=lb), tb, fmt_="0.000000", fill=COLOR_GREEN_LIGHT)
        L_abs = f"${tLB}${data0}:${tLB}${data0 + T - 1}"

    (op1, x1, _), (op2, x2, _), (_, a3, b3) = case_cells["menor"], case_cells["mayor"], case_cells["rango"]
    for j in range(n_rows):
        rr = data0 + j
        i = j + 1
        fill = COLOR_ZEBRA if i % 2 == 0 else COLOR_WHITE
        _style_cell(ws.cell(row=rr, column=1, value=i), tb, fill=fill)
        if group:
            s_i, e_i = j * group + 1, (j + 1) * group
            _style_cell(ws.cell(row=rr, column=2, value=_r(s_i) if group == 1 else f"{_r(s_i)} … {_r(e_i)}"), tb, fill=fill)
            rng = f"INDEX({R_abs},{s_i}):INDEX({R_abs},{e_i})"
            if dist == "normal":
                aux = f"=SUM({rng})"
                xf = f"={P['mu']}+{P['sigma']}*(C{rr}-6)"
                aux_fmt = "0.000000"
            elif dist == "erlang":
                aux = f"=PRODUCT({rng})"
                xf = f'=IF(C{rr}<=0,"",-(1/{P["lam"]})*LN(C{rr}))'
                aux_fmt = "0.000000E+00"
            else:
                aux = f'=COUNTIF({rng},"<"&{P["p"]})'
                xf = f"=C{rr}"
                aux_fmt = "0"
            _style_cell(ws.cell(row=rr, column=3, value=aux), tb, fmt_=aux_fmt, fill=fill)
            x_fmt = "0" if meta["discrete"] else "0.000000"
            _style_cell(ws.cell(row=rr, column=4, value=xf), tb, bold=True, fmt_=x_fmt, fill=COLOR_GREEN_LIGHT, color=COLOR_GREEN_DARK)
        elif mult:
            _style_cell(ws.cell(row=rr, column=2, value=f"={rref(i)}"), tb, fmt_="0.000000", fill=fill)
            if j == 0:
                cnt, prod = 1, f"=B{rr}"
            else:
                cnt = f"=IF(ISNUMBER(E{rr - 1}),1,C{rr - 1}+1)"
                prod = f"=IF(ISNUMBER(E{rr - 1}),B{rr},D{rr - 1}*B{rr})"
            _style_cell(ws.cell(row=rr, column=3, value=cnt), tb, fmt_="0", fill=fill)
            _style_cell(ws.cell(row=rr, column=4, value=prod), tb, fmt_="0.000000E+00", fill=fill)
            _style_cell(ws.cell(row=rr, column=5, value=f'=IF(D{rr}<{lim_cell},C{rr}-1,"")'), tb, bold=True, fmt_="0",
                        fill=COLOR_GREEN_LIGHT, color=COLOR_GREEN_DARK)
        else:
            _style_cell(ws.cell(row=rr, column=2, value=f"={rref(i)}"), tb, fmt_="0.000000", fill=fill)
            if dist == "uniforme":
                xf = f"={P['a']}+({P['b']}-{P['a']})*B{rr}"
            elif dist == "normal":
                xf = f'=IF(OR(B{rr}<=0,B{rr}>=1),"",_xlfn.NORM.INV(B{rr},{P["mu"]},{P["sigma"]}))'
            elif dist == "erlang":
                xf = f"=IF(B{rr}<=0,0,_xlfn.GAMMA.INV(B{rr},{P['k']},1/{P['lam']}))"
            else:
                xf = f"=MATCH(B{rr},{L_abs},1)-1"
            x_fmt = "0" if meta["discrete"] else "0.000000"
            _style_cell(ws.cell(row=rr, column=3, value=xf), tb, bold=True, fmt_=x_fmt, fill=COLOR_GREEN_LIGHT, color=COLOR_GREEN_DARK)

        X = f"{xl}{rr}"
        ind = [
            f'=IF(ISNUMBER({X}),IF(IF({op1}="<",{X}<{x1},{X}<={x1}),1,0),"")',
            f'=IF(ISNUMBER({X}),IF(IF({op2}=">",{X}>{x2},{X}>={x2}),1,0),"")',
            f'=IF(ISNUMBER({X}),IF(AND({X}>={a3},{X}<={b3}),1,0),"")',
        ]
        for k_, f_ in enumerate(ind):
            _style_cell(ws.cell(row=rr, column=first_ind_col + k_, value=f_), tb, fmt_="0", fill=fill)

    if n_rows == 0:
        cell = ws.cell(row=data0, column=1, value=f"No hay suficientes rᵢ para formar un solo valor con este método ({r_per_value} rᵢ por valor).")
        cell.font = Font(name="Arial", size=10, italic=True, color="B91C1C")

    # Anchos de columna
    widths = {"A": 44, "B": 16, "C": 18, "D": 18, "E": 20, "F": 20, "G": 20, "H": 16, "I": 22, "J": 18,
              "K": 16, "L": 22, "M": 16}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    info["params"] = ", ".join(
        f"{lab.split('(')[0].strip()} = {fmt(int(p[k]) if k in ('k', 'n') else p[k])}" for k, lab, _ in rows_def)
    return info


# -------------------------------------------------------------------------
# Hoja resumen
# -------------------------------------------------------------------------
def _build_summary(ws, infos: List[Dict[str, Any]], n: int, alpha: float):
    tb = create_thin_border()
    _title_banner(ws, "VALIDACIÓN DE LAS CONVERSIONES ESTADÍSTICAS — RESUMEN", span_cols=7)
    _objective_row(ws, 3, (
        f"Se tomaron los {n} números pseudoaleatorios rᵢ generados y se convirtieron a 5 distribuciones. En cada hoja V1–V5 se "
        "cuentan los casos 'menor que', 'mayor que' y 'entre a y b' y se comparan con la probabilidad teórica. "
        f"Nivel de significancia inicial α = {alpha:.2%} (editable en cada hoja)."), span_cols=7)
    ws.row_dimensions[3].height = 44
    _table_header(ws, 5, ["#", "Distribución", "Parámetros", "Método de conversión", "n (valores convertidos)",
                          "Casos validados (de 3)", "Decisión"], tb)
    ws.row_dimensions[5].height = 30
    first = 6
    for i, inf in enumerate(infos, start=1):
        r = first + i - 1
        sh = f"'{inf['sheet']}'"
        _style_cell(ws.cell(row=r, column=1, value=i), tb)
        _style_cell(ws.cell(row=r, column=2, value=inf["name"]), tb, bold=True, align="left")
        if inf.get("error"):
            _style_cell(ws.cell(row=r, column=3, value=inf["error"]), tb, align="left")
            _style_cell(ws.cell(row=r, column=4, value=inf["method"]), tb, align="left")
            _style_cell(ws.cell(row=r, column=5, value="—"), tb)
            _style_cell(ws.cell(row=r, column=6, value="—"), tb)
            _style_cell(ws.cell(row=r, column=7, value="✗ PARÁMETROS INVÁLIDOS"), tb, bold=True)
            continue
        _style_cell(ws.cell(row=r, column=3, value=inf["params"]), tb, align="left")
        _style_cell(ws.cell(row=r, column=4, value=inf["method"]), tb, align="left")
        _style_cell(ws.cell(row=r, column=5, value=f"={sh}!{inf['n_cell']}"), tb, fmt_="0")
        _style_cell(ws.cell(row=r, column=6, value=f"={sh}!{inf['ok_cell']}"), tb, bold=True, fmt_="0")
        dec = ws.cell(row=r, column=7, value=f'=IF(F{r}=3,"✓ VALIDADA","✗ NO VALIDADA")')
        _style_cell(dec, tb, bold=True, fill=COLOR_PURPLE, color="FFFFFF")
    last = first + len(infos) - 1
    tr = last + 2
    lbl = ws.cell(row=tr, column=2, value="Conversiones validadas:")
    lbl.font = Font(name="Arial", size=11, bold=True, color=COLOR_PURPLE_DARK)
    lbl.alignment = Alignment(horizontal="right", vertical="center")
    tot = ws.cell(row=tr, column=3, value=f'=COUNTIF(G{first}:G{last},"✓ VALIDADA")&" de {len(infos)}"')
    _style_cell(tot, tb, bold=True, fill=COLOR_GREEN_LIGHT, color=COLOR_GREEN_DARK)
    for col, w in {"A": 6, "B": 28, "C": 30, "D": 44, "E": 18, "F": 18, "G": 20}.items():
        ws.column_dimensions[col].width = w


def build_distribution_sheets(wb, ws1_title: str, ri_col_letter: str, data_start_row: int, n: int,
                              alpha: float, dist_config: Optional[Dict[str, Any]] = None):
    cfg = normalize_config(dist_config)
    ws_sum = wb.create_sheet(title=SUMMARY_SHEET)
    ws_sum.views.sheetView[0].showGridLines = True
    infos = [
        _build_dist_sheet(wb, i + 1, d, cfg[d], ws1_title, ri_col_letter, data_start_row, n, alpha)
        for i, d in enumerate(DIST_ORDER)
    ]
    _build_summary(ws_sum, infos, n, alpha)
